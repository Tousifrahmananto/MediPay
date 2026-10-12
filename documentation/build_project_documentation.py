from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"G:\CSE391 PROJECT")
OUT = ROOT / "documentation" / "CSE391_Tousif_Rahman_Anto_Hospital_Billing_SaaS.docx"
LOGO = ROOT / "documentation" / "bracu_logo.png"

# standard_business_brief preset with editorial_cover opening pattern.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6B7A"
LIGHT = "F2F4F7"
BORDER = "CCD5E0"
WHITE = "FFFFFF"


def set_run_font(run, size=11, color="000000", bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_fixed_table_geometry(table, widths_dxa):
    assert sum(widths_dxa) == 9360
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=10, color=INK, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(text))
            set_run_font(run, size=9.5)
    set_fixed_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def create_numbering(doc, kind="bullet"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_lead=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_callout(doc, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.05
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "E8EEF5")
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "5")
        border.set(qn("w:color"), BORDER)
        borders.append(border)
    p_pr.append(borders)
    set_run_font(p.add_run(f"{title}: "), size=10.5, color=INK, bold=True)
    set_run_font(p.add_run(text), size=10)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("CSE391 PROJECT DOCUMENTATION"), size=8.5, color=MUTED, bold=True)
    hp.add_run("\t")
    set_run_font(hp.add_run("Hospital Billing SaaS"), size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])
    add_page_number(section.first_page_footer.paragraphs[0])

    bullet_id = create_numbering(doc, "bullet")
    decimal_id = create_numbering(doc, "decimal")

    # Editorial cover pattern.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(12)
    if LOGO.exists():
        picture = p.add_run().add_picture(str(LOGO), width=Inches(0.9))
        doc_pr = picture._inline.docPr
        doc_pr.set("title", "BRAC University logo")
        doc_pr.set("descr", "BRAC University logo")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("BRAC UNIVERSITY"), size=14, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    set_run_font(p.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING"), size=10.5, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("PROJECT DOCUMENTATION"), size=10, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run("Hospital Billing SaaS"), size=28, color=INK, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40)
    set_run_font(p.add_run("A MERN-based healthcare finance and operations portal"), size=13.5, color=DARK_BLUE, italic=True)
    for label, value in (
        ("Course", "CSE391"),
        ("Prepared by", "Tousif Rahman Anto"),
        ("Email", "tousif.rahman.anto@g.bracu.ac.bd"),
        ("Submission", "Summer 2026"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(f"{label}: "), size=11, color=MUTED, bold=True)
        set_run_font(p.add_run(value), size=11, color=INK)

    doc.add_page_break()

    add_heading(doc, "1. Project Overview", 1)
    add_body(doc, "Hospital Billing SaaS is a full-stack healthcare finance platform designed for hospitals and clinics. It centralizes patient records, departmental services, charge capture, invoice generation, payment collection, appointments, documents, reports, contracts, claims, notifications and audit history in one role-based portal.")
    add_body(doc, "The system follows a multi-tenant SaaS model. Each healthcare organization owns its operational data, while a centralized super administrator can review organizations and monitor the platform. Department-scoped staff see only the functions and data permitted by their role and assigned permissions.")
    add_callout(doc, "Core objective", "Replace fragmented billing and administrative work with a traceable workflow: create the hospital structure, register a patient, record an encounter and charges, issue a readable invoice, collect payment, and retain the complete operational history.")

    add_heading(doc, "1.1 Problem Statement", 2)
    add_body(doc, "Manual hospital billing is slow and error-prone when patient details, service prices, departmental charges and payments are recorded in separate places. It is difficult to identify outstanding balances, produce consistent invoices, restrict staff access and investigate changes. This project addresses those problems through structured data, configurable service pricing, automatic calculations, immutable invoice snapshots and auditable actions.")

    add_heading(doc, "1.2 Project Objectives", 2)
    objectives = [
        "Provide secure registration, login, password recovery and role-based authorization.",
        "Organize hospital staff under departments with clear roles and permissions.",
        "Capture patient, encounter, service and charge information without duplicate manual entry.",
        "Generate itemized invoices with automatic quantity/day calculations, PDF output and email delivery.",
        "Support academic payment testing through the SSLCOMMERZ sandbox and track refunds and reconciliation.",
        "Give administrators dashboards, reports, notifications and audit history for oversight.",
    ]
    for item in objectives:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "2. Implemented Project Features", 1)
    features = [
        ("Authentication and access control: ", "organization registration, login, JWT sessions, logout, temporary passwords, forgot/reset password, roles and granular permissions."),
        ("Organization and staff: ", "organization profile and status, staff creation, department assignment, role/permission setup, reassignment and staff suspension."),
        ("Patients and departments: ", "patient records, department catalog, service catalog, encounters and department-level charge capture."),
        ("Dashboard: ", "date filters, monthly billing and collection summaries, outstanding balances, charts, invoice-status totals and operational statistics."),
        ("Appointments: ", "availability checking, booking, calendar view, rescheduling, cancellation, approval/rejection and reminder delivery."),
        ("Documents and reports: ", "file upload/download, categorization, rename/delete, report attachments, search/type/date filters and CSV export."),
        ("Contracts and claims: ", "contract upload/download, review decisions, drawn electronic signature, signed PDF generation and claim status workflow."),
        ("Invoices and payments: ", "guided charge selection, quantities and hospital-stay days, patient contact snapshots, automatic totals, PDF/email, sandbox checkout, refunds and reconciliation."),
        ("Messaging and notifications: ", "conversation history, replies, notification listing and read status."),
        ("Super administration and audit: ", "cross-organization oversight, organization review, staff controls, platform statistics and CSV audit export."),
    ]
    for lead, body in features:
        add_list_item(doc, lead + body, bullet_id, bold_lead=lead)

    add_heading(doc, "3. Technology and System Architecture", 1)
    add_body(doc, "The application uses the MERN stack with TypeScript on both the client and server. The browser client communicates with a REST API; the API applies authentication, validation and business rules before reading or writing MongoDB Atlas.")
    add_table(doc, ["Layer", "Technology", "Purpose"], [
        ("Frontend", "React, Vite, TypeScript", "Role-based portal, forms, dashboards, calendar and payment initiation"),
        ("Backend", "Node.js, Express, TypeScript", "REST endpoints, authorization, validation and workflow logic"),
        ("Database", "MongoDB Atlas, Mongoose", "Multi-tenant operational and financial records"),
        ("Security", "JWT, bcrypt", "Session authentication and password hashing"),
        ("Documents", "PDFKit and file APIs", "Invoices, receipts, contracts, reports and downloads"),
        ("Email", "Resend", "Password reset, invoice delivery and reminders"),
        ("Payment", "SSLCOMMERZ sandbox", "Hosted academic checkout and callback verification"),
        ("Hosting", "Vercel and Render", "Frontend and backend deployment"),
    ], [1600, 2600, 5160])

    add_heading(doc, "3.1 Main Data Relationships", 2)
    relationships = [
        "An organization owns users, departments, patients, appointments, documents, contracts, reports and financial records.",
        "A provider staff account may reference one active department belonging to the same organization.",
        "A patient can have multiple encounters; an encounter can collect charges from one or more hospital services.",
        "An invoice stores immutable line-item and patient-contact snapshots so later catalog edits do not change an issued bill.",
        "Payments apply to invoices; refunds reference settled payments; reconciliation compares gateway and internal transaction states.",
        "Privileged mutations create audit events containing actor, action, target and timestamp information.",
    ]
    for item in relationships:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "4. User Roles and Operational Workflow", 1)
    add_table(doc, ["Role", "Main Responsibilities"], [
        ("Super Administrator", "Reviews the whole platform, organizations, staff, documents, appointments, invoices, reports and contracts; controls organization/staff status."),
        ("Organization Owner / Admin", "Configures the hospital, creates departments and services, manages staff, approves workflows and oversees billing."),
        ("Department Staff", "Works within assigned permissions and department scope; manages permitted patients, appointments, charges, documents and reports."),
        ("Cashier / Billing Staff", "Builds invoices from charge items, verifies patient contact data, sends invoices and records payment/refund activity."),
        ("Patient / Public Payer", "Receives the invoice or payment link and completes hosted sandbox checkout without accessing internal administration."),
    ], [2200, 7160])

    add_heading(doc, "4.1 Organization and Staff Setup", 2)
    setup_steps = [
        "The organization owner registers the healthcare organization and signs in.",
        "An administrator creates departments and activates the services offered by each department.",
        "The administrator creates a staff account with name, email, phone, employee number, role and permissions.",
        "Provider staff are assigned to an active department; organization-wide administrators may remain unassigned.",
        "The API verifies that the chosen department belongs to the same organization before saving the staff record.",
        "After login, the portal menu and API operations are limited by the staff member's role, permissions and department scope.",
    ]
    for item in setup_steps:
        add_list_item(doc, item, decimal_id)

    add_heading(doc, "4.2 Patient-to-Payment Workflow", 2)
    billing_steps = [
        "Create or select a patient and record the current encounter or hospital stay.",
        "Select catalog services or add authorized charges. Quantities and number of stay days multiply automatically by the configured unit price.",
        "Review the patient name, email, phone number, line descriptions, taxes/discounts, total and due date.",
        "Issue the invoice. The server stores a permanent snapshot, generates a PDF and attempts delivery to the patient's email address.",
        "The payer opens the payment link and chooses an available SSLCOMMERZ sandbox method such as card or mobile banking.",
        "The gateway redirects to success, failure or cancellation callbacks. The server validates the transaction before marking the invoice paid.",
        "Authorized staff can later review payment history, reconcile transactions, issue eligible refunds and inspect the audit trail.",
    ]
    for item in billing_steps:
        add_list_item(doc, item, decimal_id)

    add_heading(doc, "5. Payment, Deployment and Security", 1)
    add_heading(doc, "5.1 Academic Payment Integration", 2)
    add_body(doc, "The project uses the SSLCOMMERZ sandbox because it is an academic implementation. The backend creates the checkout session, supplies transaction and callback information, and receives gateway responses. Payment status is never trusted only from a browser redirect; the backend validates the transaction before updating the invoice.")
    add_body(doc, "The Bangla QR option is presented only when supported by the configured merchant. Real production processing would require merchant onboarding, live credentials and compliance approval, which are outside the academic scope.")

    add_heading(doc, "5.2 Deployment", 2)
    add_table(doc, ["Component", "Platform", "Public Address"], [
        ("Frontend", "Vercel", "https://saas-project-client-neon.vercel.app/"),
        ("Backend API", "Render", "https://hospital-billing-api-1669.onrender.com"),
        ("Database", "MongoDB Atlas", "Managed cloud database; connection details remain private"),
        ("Source code", "GitHub", "https://github.com/Tousifrahmananto/SAAS_PROJECT"),
    ], [1800, 1800, 5760])

    add_heading(doc, "5.3 Security and Integrity Controls", 2)
    security = [
        "Passwords are stored as bcrypt hashes; reset tokens expire and are not stored as plain-text passwords.",
        "JWT authentication and server-side role/permission checks protect private routes.",
        "Organization ownership checks prevent cross-hospital access to records.",
        "Financial values use Decimal128, while issued invoice items remain immutable snapshots.",
        "Payment settlement is designed to be transactional and idempotent to reduce duplicate processing.",
        "Sensitive database, email, JWT and gateway credentials are stored only as backend environment variables.",
        "Privileged actions are written to the audit log for accountability.",
    ]
    for item in security:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "6. Testing and Verification", 1)
    add_body(doc, "The project was verified through automated tests and production builds after implementation and bug-fixing. Client and server behavior were checked independently, including the recently added administrative and workflow modules.")
    add_table(doc, ["Verification", "Result"], [
        ("Client type check", "Passed"),
        ("Server type check", "Passed"),
        ("Client automated tests", "10 passed"),
        ("Server automated tests", "20 passed"),
        ("Production build", "Passed"),
    ], [4680, 4680])

    add_heading(doc, "6.1 Important Validation Scenarios", 2)
    validations = [
        "Duplicate and invalid records are rejected by schema and API validation.",
        "Staff cannot be assigned to an inactive or foreign organization department.",
        "Appointment availability detects time conflicts before booking or rescheduling.",
        "Invoice totals are calculated from selected items, quantity/day values and stored prices.",
        "Payment callbacks are verified before financial status changes are accepted.",
        "File rename/delete/download actions remain permission-controlled.",
    ]
    for item in validations:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "7. Online Resources Used", 1)
    resources = [
        "MDN Web Docs and official React, Vite, Node.js, Express, TypeScript, MongoDB and Mongoose documentation.",
        "SSLCOMMERZ developer documentation for sandbox session creation, callbacks and validation.",
        "Vercel and Render documentation for frontend/backend deployment and environment configuration.",
        "Resend documentation for transactional email delivery.",
        "GitHub and Stack Overflow discussions were consulted for implementation-specific troubleshooting; final behavior was verified against official documentation and local tests.",
    ]
    for item in resources:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "8. Future Enhancements", 1)
    enhancements = [
        "Add optional multi-factor authentication and a dedicated editor for existing staff permissions.",
        "Add notification mark-unread controls and conversation closing/archiving.",
        "Run automatic reminder jobs through a persistent scheduler rather than relying only on the explicit reminder action.",
        "Add richer audit-log portal filters and signed PDF audit reports.",
        "Integrate an approved production payment merchant and production Bangla QR when the organization is eligible.",
        "Move large attachments to managed object storage and add retention/backup policies.",
        "Extend analytics with expenses, profitability, predictive revenue trends and department performance comparisons.",
        "Provide a patient-facing mobile application and real-time messaging where required.",
    ]
    for item in enhancements:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "9. Conclusion", 1)
    add_body(doc, "Hospital Billing SaaS demonstrates a complete MERN workflow for healthcare administration and billing. It connects hospital structure, staff authorization, patient services, invoicing, academic payment processing and administrative oversight while preserving data separation and auditability. The completed system meets the core project objective and leaves a clear path from academic sandbox deployment to a production-ready healthcare finance platform.")

    core = doc.core_properties
    core.title = "Hospital Billing SaaS - CSE391 Project Documentation"
    core.subject = "Full-stack healthcare SaaS project report"
    core.author = "Tousif Rahman Anto"
    core.keywords = "CSE391, MERN, Hospital Billing, SaaS, MongoDB, React, Express"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
